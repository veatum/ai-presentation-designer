from __future__ import annotations
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor


def build_default_template(out_path: str | Path) -> Path:
    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)

    # Keep the fallback intentionally original and minimal: it is not a copy of any third-party deck.
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    title = slide.shapes.add_textbox(Inches(0.85), Inches(0.7), Inches(11.6), Inches(1.0))
    p = title.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = "Presentation title"
    r.font.name = "Arial"
    r.font.size = Pt(30)
    r.font.bold = True
    r.font.color.rgb = RGBColor(32, 35, 43)

    subtitle = slide.shapes.add_textbox(Inches(0.85), Inches(1.8), Inches(11.0), Inches(0.7))
    p = subtitle.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = "A clean editable starting point"
    r.font.name = "Arial"
    r.font.size = Pt(17)
    r.font.color.rgb = RGBColor(95, 101, 115)

    # Add a simple content slide so the renderer has a known body pattern.
    slide = prs.slides.add_slide(blank)
    title = slide.shapes.add_textbox(Inches(0.85), Inches(0.55), Inches(11.6), Inches(0.8))
    p = title.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = "Section title"
    r.font.name = "Arial"
    r.font.size = Pt(28)
    r.font.bold = True
    body = slide.shapes.add_textbox(Inches(0.85), Inches(1.7), Inches(11.6), Inches(4.7))
    body.text_frame.word_wrap = True
    p = body.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = "Editable content"
    r.font.name = "Arial"
    r.font.size = Pt(18)
    r.font.color.rgb = RGBColor(45, 49, 58)

    prs.save(str(out_path))
    return Path(out_path)
