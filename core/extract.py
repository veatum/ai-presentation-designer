from pathlib import Path
import json
from pypdf import PdfReader
from docx import Document

MAX_TEXT=55000

def extract_text(path):
    p=Path(path); ext=p.suffix.lower()
    if ext in {'.txt','.md','.csv','.json'}: return p.read_text(encoding='utf-8-sig')
    if ext=='.pdf':
        reader=PdfReader(str(p))
        if len(reader.pages)>100: raise ValueError('Не более 100 страниц PDF.')
        return '\n'.join(page.extract_text() or '' for page in reader.pages)
    if ext=='.docx':
        doc=Document(str(p)); parts=[p.text for p in doc.paragraphs if p.text.strip()]
        parts += [' | '.join(c.text for c in row.cells) for table in doc.tables for row in table.rows]
        return '\n'.join(parts)
    raise ValueError('Неподдерживаемый формат источника')

def extract_source_records(paths,names=None):
    records=[]; total=0
    for i,path in enumerate(paths,1):
        text=extract_text(path).strip(); total+=len(text)
        if not text: raise ValueError('В файле нет извлекаемого текста; для сканов требуется OCR.')
        if total>MAX_TEXT: raise ValueError('Материалы превышают 55 000 символов. Разделите файлы; текст не обрезан.')
        record={'id':f'U{i}','kind':'upload','title':names[i-1] if names else Path(path).name,'text':text,'url':'','accessed_at':'','verified':False,'verification':'user supplied; not independently verified'}
        if Path(path).suffix.lower()=='.json':
            document=json.loads(text)
            if isinstance(document,dict) and isinstance(document.get('text'),str) and isinstance(document.get('bibliography'),dict):
                record['text']=document['text']
                allowed={'reference_kind','author','title','publisher','year','pages','journal','volume','issue','doi','organization','document_date','document_number','place'}
                record.update({k:str(v) for k,v in document['bibliography'].items() if k in allowed and isinstance(v,(str,int))})
                record['metadata_origin']='user-supplied; verify bibliographic details before publication'
        records.append(record)
    return records

def extract_sources(paths):
    return '\n\n'.join('### '+s['title']+'\n'+s['text'] for s in extract_source_records(paths))
