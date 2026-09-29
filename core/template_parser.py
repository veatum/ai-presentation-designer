from __future__ import annotations
import collections
import math
from pathlib import Path
from typing import Any
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

EMU_PER_INCH = 914400


def _rgb(fill) -> str | None:
    try:
        if fill.type is not None and fill.fore_color.type is not None:
            rgb = fill.fore_color.rgb
            if rgb:
                return str(rgb)
    except Exception:
        return None
    return None


def _font_from_shape(shape) -> str | None:
    try:
        for p in getattr(shape.text_frame, "paragraphs", []):
            for run in p.runs:
                if run.font.name:
                    return run.font.name
    except Exception:
        return None
    return None


def _placeholder_kind(ph) -> str:
    try:
        name = str(ph.placeholder_format.type).lower()
    except Exception:
        return "unknown"
    if "subtitle" in name:
        return "subtitle"
    if "title" in name or "centertitle" in name:
        return "title"
    if "body" in name or "content" in name or "text" in name or "object" in name:
        return "body"
    if "picture" in name:
        return "picture"
    return "other"


def _classify_slide(slide) -> str:
    texts = []
    placeholders = []
    for shape in slide.shapes:
        if getattr(shape, "has_text_frame", False):
            t = shape.text.strip()
            if t:
                texts.append(t)
        if getattr(shape, "is_placeholder", False):
            placeholders.append(_placeholder_kind(shape))
    if placeholders.count("body") >= 2:
        return "two_column"
    if "title" in placeholders and "body" in placeholders:
        return "content"
    if any("table" in str(getattr(shape, "shape_type", "")).lower() for shape in slide.shapes):
        return "table"
    if len(texts) <= 2:
        return "section"
    return "content"


def parse_template(path: str | Path) -> dict[str, Any]:
    prs = Presentation(str(path))
    fonts = collections.Counter()
    colors = collections.Counter()
    slides = []
    layout_summaries = []

    for index, slide in enumerate(prs.slides, start=1):
        shapes = []
        for shape in slide.shapes:
            item = {
                "type": str(getattr(shape, "shape_type", "unknown")),
                "left": round(shape.left / EMU_PER_INCH, 3),
                "top": round(shape.top / EMU_PER_INCH, 3),
                "width": round(shape.width / EMU_PER_INCH, 3),
                "height": round(shape.height / EMU_PER_INCH, 3),
            }
            if getattr(shape, "has_text_frame", False):
                text = shape.text.strip()
                item["text_sample"] = text[:180]
                item["is_placeholder"] = bool(getattr(shape, "is_placeholder", False))
                if getattr(shape, "is_placeholder", False):
                    try:
                        item["placeholder_kind"] = _placeholder_kind(shape)
                    except Exception:
                        pass
                font = _font_from_shape(shape)
                if font:
                    fonts[font] += 1
            try:
                color = _rgb(shape.fill)
                if color:
                    colors[color] += 1
            except Exception:
                pass
            shapes.append(item)

        kind = _classify_slide(slide)
        slides.append({"index": index, "kind": kind, "shape_count": len(shapes), "shapes": shapes})

    for idx, layout in enumerate(prs.slide_layouts):
        phs = []
        for ph in layout.placeholders:
            if any(v is None for v in (ph.left,ph.top,ph.width,ph.height)):
                continue
            phs.append({
                "kind": _placeholder_kind(ph),
                "left": round(ph.left / EMU_PER_INCH, 3),
                "top": round(ph.top / EMU_PER_INCH, 3),
                "width": round(ph.width / EMU_PER_INCH, 3),
                "height": round(ph.height / EMU_PER_INCH, 3),
            })
        layout_summaries.append({"index": idx, "name": layout.name, "placeholders": phs})

    return {
        "path": str(path),
        "slide_size": {"width_in": round(prs.slide_width / EMU_PER_INCH, 3), "height_in": round(prs.slide_height / EMU_PER_INCH, 3)},
        "slide_count": len(prs.slides),
        "design_tokens": {
            "fonts": fonts.most_common(5),
            "colors": colors.most_common(10),
        },
        "slides": slides,
        "layouts": layout_summaries,
    }
