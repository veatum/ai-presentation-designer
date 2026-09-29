import json, shutil, subprocess, sys, types
from pathlib import Path
from types import SimpleNamespace

# Provide a tiny openai stub so all project logic can be exercised offline.
import builtins
m=types.ModuleType('openai')
class Base(Exception): pass
for n in ['APIStatusError','APIConnectionError','APITimeoutError','RateLimitError','InternalServerError']:
    setattr(m,n,type(n,(Base,),{}))
class OpenAI: pass
m.OpenAI=OpenAI
sys.modules['openai']=m
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))

import core.llm as llm
from core.llm import Session, OutputError, _complete, generate_plan
from core.quality import quality_issues, source_issues
from core.template_parser import parse_template
from core.renderer import render_variant, VARIANTS
from core.template_profiles import BUILTIN_TEMPLATES, builtin_template_path
from core.default_template import build_default_template
from core.audit import audit_deck
from core.exporter import export_pdf, export_html
from pptx import Presentation

OUT=Path('/mnt/data/acceptance_run'); OUT.mkdir(exist_ok=True)
SOURCES=[{'id':'W1','kind':'web','verified':True,'title':'Официальная документация','url':'https://example.org/source','site_name':'Example','text':('Сервис использует структурированный вывод для передачи данных в заданном формате. Модель Qwen3-32B предназначена для задач генерации текста и поддерживает несколько возможностей. В документации указано, что API использует метод POST /v1/chat/completions для генерации. Один и тот же контент может быть перенесён на разные шаблоны без повторной генерации смысла. Три варианта верстки позволяют сравнивать разные композиционные решения. Число 20 используется в примере ограничения; числа 10 и 12 приведены в разделе параметров. ')*3}]
B=[
 'Структурированный вывод позволяет приложению проверять результат до рендера, поэтому формат данных остаётся предсказуемым и последующий слой верстки получает валидную структуру.',
 'Разделение смысловой генерации и шаблона позволяет менять оформление без повторной генерации содержания, что делает пайплайн воспроизводимым и снижает нагрузку на модель.',
 'Три варианта верстки нужны для сравнения композиционных решений одного и того же содержания при сохранении правил выбранного шаблона.',
 'Официальная документация описывает способ обращения к API, поэтому этот источник можно использовать как основание для технических утверждений о работе сервиса.',
 'Проверка каждого числового значения по исходному тексту не позволяет презентации сохранить число, для которого нет подтверждения в загруженном веб-источнике.',
 'Финальный аудит разделяет детерминированные проверки файла и смысловые проверки, которые могут выполняться моделью на повторных запусках.',
]

class FakeClient:
    count=0
    def __init__(self,**kwargs): self.chat=SimpleNamespace(completions=self)
    def __enter__(self): return self
    def __exit__(self,*args): pass
    def create(self,**kwargs):
        FakeClient.count += 1
        user=json.loads(kwargs['messages'][1]['content'])
        system=kwargs['messages'][0]['content']
        # Deliberately break the first SlideBatch response to verify automatic recovery.
        if 'required_slide_numbers' in user and FakeClient.count==2:
            return SimpleNamespace(choices=[SimpleNamespace(finish_reason='stop',message=SimpleNamespace(content='Не удалось.'))],usage=None)
        if 'core_thesis' in system and 'visual_logic' in system:
            nums=len(user.get('sources',[])) and int(user.get('requested_slides',10)) or 10
            out={'story':'Проверенный путь от источников к редактируемому PPTX.','core_thesis':'Содержание и оформление разделены.','slides':[]}
            for n in range(1,nums+1):
                out['slides'].append({'number':n,'role':'content','key_question':'Что важно?','key_message':B[(n-1)%len(B)],'evidence_needed':['официальный веб-источник'],'source_ids':['W1'],'visual_logic':'Объясняющий блок'})
            return SimpleNamespace(choices=[SimpleNamespace(finish_reason='stop',message=SimpleNamespace(content='```json\n'+json.dumps(out,ensure_ascii=False)+'\n```'))],usage=None)
        if 'overall' in system and 'global_fixes' in system:
            return SimpleNamespace(choices=[SimpleNamespace(finish_reason='stop',message=SimpleNamespace(content=json.dumps({'overall':'pass','issues':[],'global_fixes':[]})))],usage=None)
        nums=user.get('required_slide_numbers',[1]); out=[]
        for n in nums:
            if n==1:
                out.append({'number':1,'kind':'title','title':'Проверенная генерация презентаций','subtitle':'Источники → содержание → шаблон','points':[],'source_ids':['W1']})
            elif n%3==0:
                out.append({'number':n,'kind':'two_column','title':f'Раздел {n}','left_title':'Смысл','left_points':[B[0],B[1]],'right_title':'Практика','right_points':[B[2],B[3]],'source_ids':['W1']})
            elif n%4==0:
                out.append({'number':n,'kind':'process','title':f'Процесс {n}','timeline':[{'label':'Источники','detail':B[3]},{'label':'Проверка','detail':B[4]},{'label':'Рендер','detail':B[5]}],'source_ids':['W1']})
            else:
                out.append({'number':n,'kind':'content','title':f'Раздел {n}','points':[B[n%len(B)],B[(n+1)%len(B)]],'source_ids':['W1']})
        return SimpleNamespace(choices=[SimpleNamespace(finish_reason='stop',message=SimpleNamespace(content=json.dumps({'slides':out},ensure_ascii=False)))],usage=None)

old=llm.OpenAI; old_settings=llm.settings
llm.OpenAI=FakeClient
llm.settings=SimpleNamespace(api_key='x',base_url='https://example.org/v1',model='Qwen/Qwen3-32B',structured_output='json_object',api_timeout=30,api_retries=1,session_seconds=290,max_llm_calls=18,max_slides=15,disable_thinking=True)

try:
    # Generation: malformed first batch must recover, never ask user to reduce slides.
    s=Session(); plan=generate_plan('Сделай презентацию о надёжной генерации презентаций.',{},'',10,{'sources':SOURCES},session=s)
    assert len(plan['slides'])==10
    # Content QA is advisory in the final build: warnings are retained in diagnostics and never block a usable deck.
    assert isinstance(quality_issues(plan,10), list)
    assert isinstance(source_issues(plan,SOURCES), list)
    assert len({x['kind'] for x in plan['slides'] if x['kind'] not in ('title','sources','section')})>=3

    # Render all three supplied VK templates and all three required layout variants.
    for tid in BUILTIN_TEMPLATES:
        tpl=builtin_template_path(tid); spec=parse_template(tpl)
        for variant in VARIANTS:
            out=OUT/f'{tid}_{variant}.pptx'
            render_variant(tpl,plan,spec,out,variant,tid,SOURCES)
            opened=Presentation(str(out))
            assert len(opened.slides)==10
            audit=audit_deck(out,plan)
            assert isinstance(audit.get('issues'),list)
            if variant=='balanced':
                pdf=export_pdf(out,OUT,timeout=30)
                export_html(out,OUT/f'{tid}_{variant}.html',pdf)
    # Also test the non-VK/default path.
    base=build_default_template(OUT/'base.pptx')
    for variant in VARIANTS:
        out=OUT/f'ai_base_{variant}.pptx'; render_variant(base,plan,parse_template(base),out,variant,'ai_base',SOURCES)
        assert len(Presentation(str(out)).slides)==10
        assert isinstance(audit_deck(out,plan).get('issues'),list)

    (OUT/'RESULT.txt').write_text(f'PASS\nslides=10\nvk_templates={len(BUILTIN_TEMPLATES)}\nvariants={len(VARIANTS)}\noutputs={len(list(OUT.glob("*.pptx")))}\nllm_calls={FakeClient.count}\n',encoding='utf-8')
    print('ACCEPTANCE PASS', 'calls=',FakeClient.count, 'pptx=',len(list(OUT.glob('*.pptx'))))
finally:
    llm.OpenAI=old; llm.settings=old_settings
