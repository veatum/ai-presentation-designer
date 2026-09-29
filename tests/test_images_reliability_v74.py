from pathlib import Path
from PIL import Image
from pptx import Presentation

from tests.test_pipeline import sample
from core.renderer import render_variant, VARIANTS
from core.template_profiles import BUILTIN_TEMPLATES, builtin_template_path
from core.template_parser import parse_template


def test_manual_image_is_embedded_without_breaking_any_vk_variant(tmp_path):
    plan=sample()
    img=tmp_path/'topic.jpg'
    Image.new('RGB',(1600,900),(90,120,160)).save(img,'JPEG')
    info={'slide':2,'path':str(img),'title':'Тематическая иллюстрация','creator':'User','license':'User-provided','landing_url':''}
    for tid in BUILTIN_TEMPLATES:
        tpl=builtin_template_path(tid); spec=parse_template(tpl)
        for variant in VARIANTS:
            out=tmp_path/f'{tid}_{variant}.pptx'
            render_variant(tpl,plan,spec,out,variant,tid,[],[info])
            prs=Presentation(out)
            assert len(prs.slides)==len(plan['slides'])
            assert out.stat().st_size>20_000


def test_renderer_recovers_from_a_layout_exception(tmp_path, monkeypatch):
    import core.renderer as r
    plan=sample()
    tpl=builtin_template_path('vk_tech'); spec=parse_template(tpl)
    original=r._fill
    def boom(*args,**kwargs):
        raise r.LayoutError('synthetic layout failure')
    monkeypatch.setattr(r,'_fill',boom)
    out=tmp_path/'recovered.pptx'
    r.render_variant(tpl,plan,spec,out,'balanced','vk_tech',[])
    prs=Presentation(out)
    assert len(prs.slides)==len(plan['slides'])
    monkeypatch.setattr(r,'_fill',original)
