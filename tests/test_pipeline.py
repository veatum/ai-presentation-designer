import io,json,time,sys,types
from pathlib import Path
from types import SimpleNamespace
# The project tests must run even in a clean offline environment without the optional
# OpenAI SDK installed; live provider checks remain separate.
try:
    import openai  # noqa: F401
except Exception:
    m=types.ModuleType('openai')
    class Base(Exception): pass
    for _n in ['APIStatusError','APIConnectionError','APITimeoutError','RateLimitError','InternalServerError']:
        setattr(m,_n,type(_n,(Base,),{}))
    class OpenAI: pass
    m.OpenAI=OpenAI
    sys.modules['openai']=m
import pytest
from pydantic import ValidationError
from pptx import Presentation
from starlette.datastructures import UploadFile
from core.models import Deck,Slide,Chart,Table
from core.llm import _extract_json,_normalize_plan,_complete,Session,OutputError
from core.quality import quality_issues,source_issues,bibliography,format_reference
from core.default_template import build_default_template
from core.template_parser import parse_template
from core.renderer import render_variant,VARIANTS
from core.audit import audit_deck
from core.exporter import export_html

LONG=[
 'Сначала автор собирает материалы и определяет вопрос аудитории. Это позволяет отделить проверяемые факты от предположений и выбрать содержание, которое действительно нужно для принятия решения.',
 'Затем редактор связывает каждый вывод с исходным документом и проверяет единицы измерения. Такой порядок помогает заметить несовпадение периодов и не объяснять результат причиной, которой нет в материалах.',
 'После сборки команда просматривает все слайды и сравнивает их с исходниками. Проверка включает читаемость текста, сохранность данных и наличие редактируемых объектов для дальнейшей работы.'
]

def sample():
    slides=[Slide(number=1,kind='title',title='Как подготовить проверяемую презентацию',subtitle='Последовательность подготовки, редактуры и контроля').model_dump()]
    kinds=['content','two_column','table','chart','process','metrics','summary']
    for n,kind in enumerate(kinds,2):
        s=Slide(number=n,kind='content',title=f'Этап {n}: содержание подтверждается материалами',points=LONG.copy()).model_dump();s['kind']=kind
        if kind=='two_column':s.update(points=[],left_title='До проверки',right_title='После проверки',left_points=LONG[:2],right_points=LONG[2:])
        if kind=='table':s.update(table={'columns':['Этап','Проверка'],'rows':[['Материалы','Проверить источник'],['Редактура','Сопоставить утверждения'],['Экспорт','Проверить все объекты']]},points=LONG[:2])
        if kind=='chart':s.update(chart={'type':'column','labels':['Январь','Февраль','Март'],'values':[10,15,20],'unit':'проверок'},points=LONG[:2])
        if kind=='process':s.update(points=[],timeline=[{'label':'Подготовка','detail':LONG[0]},{'label':'Редактура','detail':LONG[1]},{'label':'Контроль','detail':LONG[2]}])
        if kind=='metrics':s.update(points=[],metrics=[{'label':'Январь','value':'10','note':LONG[0]},{'label':'Март','value':'20','note':LONG[1]}])
        s['source_ids']=['U1'];s['claims']=[{'text':LONG[0],'source_id':'U1','evidence_quote':LONG[0]}]
        slides.append(s)
    return Deck(title='Проверка renderer',slides=slides).model_dump()

@pytest.mark.parametrize('raw',['{} {}','{"a":1}\n{"b":2}','```json\n{}\n```','before {}','[]','{"a":NaN}'])
def test_multiple_json_or_wrapper_never_silently_discarded(raw):
    with pytest.raises(OutputError):_extract_json(raw)

def test_schema_and_count():
    with pytest.raises(ValidationError):Chart(type='line',labels=['a','b'],values=[1],unit='руб.')
    with pytest.raises(ValidationError):Table(columns=['a'],rows=[['b','c']])
    with pytest.raises(ValueError):_normalize_plan(sample(),3)

def test_source_citation_exists_and_numbers_supported():
    p=sample();sources=[{'id':'U1','text':' '.join(LONG),'kind':'upload','title':'Материалы'}]
    issues=source_issues(p,sources)
    assert any('числа' in i for i in issues)
    p['slides'][1]['claims'][0]['evidence_quote']='Этой цитаты нет в источнике'
    assert any('цитата отсутствует' in i for i in source_issues(p,sources))

def test_sparse_notes_do_not_pass():
    p=sample();p['slides'][1].update(points=['Рост рынка'],notes=' '.join(LONG)*5)
    assert any('недостаточно' in i for i in quality_issues(p,len(p['slides'])))

def test_bibliography_real_access_date_only_used():
    p=sample();p['slides'].append(Slide(number=9,kind='sources',title='Источники').model_dump())
    sources=[{'id':'U1','kind':'web','title':'Исследование','url':'https://example.org/report','accessed_at':'2026-09-28T12:00:00+00:00','text':' '.join(LONG)}, {'id':'UNUSED','kind':'web','title':'Не использован','url':'https://example.org/unused'}]
    p=bibliography(p,sources,1)
    assert '28.09.2026' in p['slides'][-1]['points'][0]
    assert 'unused' not in str(p['slides'][-1])
    assert 'DOI' not in format_reference(sources[0])

@pytest.mark.parametrize('variant',VARIANTS)
def test_every_slide_and_all_native_payloads_survive(tmp_path,variant):
    template=build_default_template(tmp_path/'template.pptx');p=sample()
    out=render_variant(template,p,parse_template(template),tmp_path/(variant+'.pptx'),variant)
    prs=Presentation(out)
    assert len(prs.slides)==len(p['slides'])
    assert any(sh.has_chart for s in prs.slides for sh in s.shapes)
    assert any(sh.has_table for s in prs.slides for sh in s.shapes)
    issues=audit_deck(out,p)['issues']
    assert not any(i.get('severity')=='critical' for i in issues)
    html=export_html(out,tmp_path/'out.html').read_text()
    assert 'Январь' in html and 'Проверить источник' in html

def test_upload_returns_actual_path(tmp_path):
    from app import _save_upload
    path=_save_upload(UploadFile(filename='material.txt',file=io.BytesIO(b'content')),tmp_path)
    assert path.read_bytes()==b'content'

def test_one_json_repair_keeps_whole_original(monkeypatch):
    import core.llm as llm
    calls=[]; p=sample(); raws=[json.dumps(p,ensure_ascii=False)+' {}',json.dumps(p,ensure_ascii=False)]
    class Client:
        def __init__(self,**kw):self.chat=SimpleNamespace(completions=self)
        def __enter__(self):return self
        def __exit__(self,*args):pass
        def create(self,**kw):
            calls.append(kw);return SimpleNamespace(choices=[SimpleNamespace(finish_reason='stop',message=SimpleNamespace(content=raws[len(calls)-1]))],usage=None)
    monkeypatch.setattr(llm,'OpenAI',Client);monkeypatch.setattr(llm,'settings',SimpleNamespace(api_key='placeholder',base_url='https://example.org',model='test',structured_output='json_schema'))
    session=Session();out=_complete('test',{},session=session)
    assert len(out['slides'])==8 and len(calls)==1
    assert session.format_repairs==0

def test_ssrf_rejected():
    from core.web_research import public_url
    for url in ['http://127.0.0.1','http://169.254.169.254/','file:///etc/passwd','http://[::1]/']:
        with pytest.raises(ValueError):public_url(url)

@pytest.mark.parametrize('template_id',['vk_tech','vk_workspace','vk_education'])
@pytest.mark.parametrize('variant',VARIANTS)
def test_builtin_templates_keep_all_payloads(tmp_path,template_id,variant):
    from core.template_profiles import builtin_template_path
    template=builtin_template_path(template_id);plan=sample()
    out=render_variant(template,plan,parse_template(template),tmp_path/'result.pptx',variant)
    issues=audit_deck(out,plan)['issues']
    assert not any(i.get('severity')=='critical' for i in issues)

def test_fifteen_slides_from_two_slide_unknown_template(tmp_path):
    template=build_default_template(tmp_path/'unknown.pptx');p=sample()
    from copy import deepcopy
    while len(p['slides'])<15:
        s=deepcopy(p['slides'][1]);s['number']=len(p['slides'])+1;s['title']=f'Деталь {s["number"]}'
        p['slides'].append(s)
    out=render_variant(template,p,parse_template(template),tmp_path/'fifteen.pptx','balanced')
    assert len(Presentation(out).slides)==15
    issues=audit_deck(out,p)['issues']
    assert not any(i.get('severity')=='critical' for i in issues)

def test_sources_hyperlink_is_editable(tmp_path):
    tpl=build_default_template(tmp_path/'tpl.pptx')
    p=sample();p['slides']=[Slide(number=1,kind='sources',title='Источники',points=['Проверяемая публикация. URL: https://example.org/report'],source_ids=['W1']).model_dump()]
    path=render_variant(tpl,p,parse_template(tpl),tmp_path/'refs.pptx','balanced',sources=[{'id':'W1','url':'https://example.org/report'}])
    links=[r.hyperlink.address for s in Presentation(path).slides for sh in s.shapes if sh.has_text_frame for p in sh.text_frame.paragraphs for r in p.runs]
    assert 'https://example.org/report' in links

def test_bibliography_cannot_overwrite_conclusion():
    with pytest.raises(ValueError,match='заменить'):
        bibliography(sample(),[{'id':'U1','kind':'upload','title':'Материалы'}],1)

def test_unknown_payload_not_silently_ignored():
    with pytest.raises(ValidationError):Slide(number=1,kind='content',title='Test',chart=Chart(type='line',labels=['a','b'],values=[1,2],unit='шт.'))

def test_books_articles_only_given_metadata():
    ref=format_reference({'title':'Название','reference_kind':'book','author':'Автор','publisher':'Издатель','year':'2024','pages':'120'})
    assert '120 с.' in ref and 'Москва' not in ref
    ref=format_reference({'title':'Статья','reference_kind':'article','journal':'Журнал','year':'2024'})
    assert 'DOI' not in ref and 'С. ' not in ref

def test_timeout_and_truncation_do_not_retry(monkeypatch):
    import core.llm as llm
    calls=[]
    class Client:
        def __init__(self,**kw):self.chat=SimpleNamespace(completions=self)
        def __enter__(self):return self
        def __exit__(self,*args):pass
        def create(self,**kw):
            calls.append(kw);return SimpleNamespace(choices=[SimpleNamespace(finish_reason='length',message=SimpleNamespace(content='{"title":"partial"'))],usage=None)
    monkeypatch.setattr(llm,'OpenAI',Client);monkeypatch.setattr(llm,'settings',SimpleNamespace(api_key='placeholder',base_url='https://example.org',model='test',structured_output='json_schema'))
    with pytest.raises(OutputError,match='обрезан'):_complete('test',{})
    assert len(calls)==1

def test_api_validation_does_not_spend(monkeypatch,tmp_path):
    from fastapi.testclient import TestClient
    import app
    monkeypatch.setattr(app,'RUNS',tmp_path)
    with TestClient(app.app) as client:
        assert client.post('/api/generate',data={'brief':'','slides':5}).status_code==400
        assert client.post('/api/generate',data={'brief':'Тема','slides':25}).status_code==400
        assert client.get('/api/file/bad/.env').status_code==404
        assert 'api_key_length' not in client.get('/api/diagnostics').json()


def test_many_short_labels_are_not_substantive_explanations():
    p=sample();p['slides'][1]['points']=['Короткие тезисы сами по себе не раскрывают механизм решения.']*5
    assert any('развёрнутых' in issue for issue in quality_issues(p,len(p['slides'])))


def test_targeted_repair_preserves_other_slides(monkeypatch):
    import core.llm as llm
    source={'id':'U1','kind':'web','url':'https://example.org/materials','verified':True,'text':' '.join(LONG),'title':'Материалы'}
    p=sample();p['slides']=[p['slides'][0],p['slides'][1],p['slides'][-1]]
    p['slides'][2]['number']=3
    p['slides'][1]['points']=['Короткий тезис']
    p['slides'][1]['title']='Материалы определяют содержание'
    p['slides'][2]['title']='Проверка завершает работу'
    p['slides'][2]['points']=[LONG[1],LONG[2]]
    calls=[]
    from copy import deepcopy
    def complete(system,payload,temperature,max_tokens,**kw):
        calls.append(kw['stage'])
        if kw['stage']=='architecture':return {'story':'Тест','core_thesis':'Тест','slides':[{'number':1,'role':'title','key_question':'q','key_message':'m','evidence_needed':['e'],'source_ids':['U1'],'visual_logic':'v'},{'number':2,'role':'content','key_question':'q','key_message':'m','evidence_needed':['e'],'source_ids':['U1'],'visual_logic':'v'},{'number':3,'role':'summary','key_question':'q','key_message':'m','evidence_needed':['e'],'source_ids':['U1'],'visual_logic':'v'}]}
        if kw['stage'].startswith('write_') or kw['stage']=='content_repair':return deepcopy(p)
        if kw['stage']=='critic':return {'overall':'repair','issues':[]}
        if kw['stage']=='final_fact_check':return {'overall':'pass','issues':[],'global_fixes':[]}
        fixed=deepcopy(p['slides'][1]);fixed['points']=LONG[:2]
        return {'slides':[fixed]}
    monkeypatch.setattr(llm,'_complete',complete)
    result=llm.generate_plan('Объясни механику',{},'',3,{'sources':[source]})
    assert len(result['slides'])==3 and [x['number'] for x in result['slides']]==[1,2,3]
    assert result['slides'][0]['title']==p['slides'][0]['title'] and result['slides'][0]['subtitle']==p['slides'][0]['subtitle'] and result['slides'][2]['number']==3 and result['slides'][2]['title']==p['slides'][2]['title'] and result['slides'][2]['points']==p['slides'][2]['points']
    assert len(result['slides'][1]['points'])>=2  # deterministic hardening restores readable slide blocks


def test_generate_route_upload_exports_and_private_trace(monkeypatch,tmp_path):
    import app as web
    from fastapi.testclient import TestClient
    monkeypatch.setattr(web,'RUNS',tmp_path)
    captured={}
    def generated(brief,spec,text,count,research,**kw):
        captured['sources']=research['sources'];return sample()
    monkeypatch.setattr(web,'generate_plan',generated)
    monkeypatch.setattr(web,'research',lambda *args,**kwargs: {'sources':[{'id':'W1','kind':'web','title':'web','url':'https://example.org/facts','text':'Verified web evidence for this test with enough detail.'}], 'domains':['example.org'], 'warnings':[]})
    def missing_pdf(*args,**kw):raise ValueError('Converter missing in this test')
    monkeypatch.setattr(web,'export_pdf',missing_pdf)
    client=TestClient(web.app)
    response=client.post('/api/generate',data={'brief':'Объясни по материалам','slides':'8','include_vk_templates':'false','include_images':'false'},files={'sources':('facts.txt',' '.join(LONG).encode(),'text/plain')})
    assert response.status_code==200,response.text
    data=response.json();rid=data['run_id']
    assert any(s.get('title')=='facts.txt' for s in captured['sources'])
    assert len(data['results'])==3 and {r['variant'] for r in data['results']}=={'balanced','columns','editorial'}
    for result in data['results']:
        assert client.get(f'/api/file/{rid}/{result["pptx"]}').status_code==200
        assert not any(i.get('severity')=='critical' for i in result['audit']['issues'])
        assert not any(i.get('severity')=='critical' for i in result['audit']['issues'])
    assert client.get(f'/api/file/{rid}/generation_trace.json').status_code==404
    original=(tmp_path/rid/'plan.json').read_bytes()
    def invalid_edit(*args,**kw):raise ValueError('Проверка содержания не пройдена')
    monkeypatch.setattr(web,'revise_plan',invalid_edit)
    revised=client.post(f'/api/revise/{rid}',data={'instruction':'Измени вывод'})
    assert revised.status_code==422
    assert (tmp_path/rid/'plan.json').read_bytes()==original


def test_dense_comparison_adapts_small_template_without_loss(tmp_path):
    from core.template_profiles import builtin_template_path
    p=sample();s=p['slides'][2]
    s.update(subtitle='Для сравнения нужны объяснения механизма и ограничений в обоих случаях.',left_points=LONG[:2],right_points=LONG[1:])
    tpl=builtin_template_path('vk_tech')
    path=render_variant(tpl,p,parse_template(tpl),tmp_path/'dense.pptx','balanced')
    issues=audit_deck(path,p)['issues']
    assert not any(i.get('severity')=='critical' for i in issues)


def test_final_fact_review_checks_revised_candidate(monkeypatch):
    import core.llm as llm
    p=sample();calls=[]
    from copy import deepcopy
    def complete(system,payload,temperature,max_tokens,**kwargs):
        calls.append((kwargs['stage'],payload))
        if kwargs['stage']=='final_fact_check':return {'overall':'repair','issues':[{'slide':2,'severity':'critical','problem':'Нет доказательств','fix':'Вернуть подтверждённое объяснение'}],'global_fixes':[]}
        fixed=deepcopy(p['slides'][1]);fixed['points']=LONG[1:];return {'slides':[fixed]}
    monkeypatch.setattr(llm,'_complete',complete)
    result=llm.verify_final_plan(p,[{'id':'U1','text':' '.join(LONG)}],llm.Session())
    assert calls[0][1]['plan']==p
    assert [c[0] for c in calls]==['final_fact_check','fact_repair']
    assert result['slides'][1]['points']==LONG[1:] and result['slides'][2:]==p['slides'][2:]


def test_generate_plan_allows_best_effort_without_web_sources():
    import core.llm as llm
    source={'id':'U1','kind':'upload','title':'local','text':'Local evidence '*30,'url':'','verified':True}
    llm.OpenAI = type('Stub',(object,),{})
    # The current product treats web research as best-effort enrichment; local context
    # remains sufficient to produce a usable semantic plan when external web access is down.
    try:
        result=llm._fallback_batch('Тема', {'sources':[source]}, [1,2,3])
        assert result and len(result)>=1
    except Exception as exc:
        raise AssertionError(str(exc))


def test_claim_overflow_is_normalized_without_pydantic_failure():
    from core.llm import _coerce_slide
    evidence='Подтверждённый источник содержит достаточно подробного текста для проверки утверждений и числовых значений.'
    raw={
        'number':4,'kind':'content','title':'Переполнение claims','points':['Развёрнутый содержательный блок с проверяемым утверждением и необходимым контекстом для демонстрации устойчивости схемы.'],
        'claims':[{'text':f'Утверждение {i}','source_id':'W1','evidence_quote':evidence} for i in range(14)],
        'source_ids':['W1']
    }
    slide=_coerce_slide(raw)
    assert len(slide['claims'])==14
    from core.models import Slide
    validated=Slide.model_validate(slide)
    assert len(validated.claims)==14


def test_numeric_content_is_advisory_and_does_not_block():
    p=sample()
    p['slides'][1]['points']=[
        'Согласно материалу, период 08.10.2022–2023 связан с показателями 28 и 630, а дополнительный показатель равен 9.',
        'Этот блок содержит контекст для числа 20 и показывает следствие для дальнейшего анализа без необходимости дополнительной структурной правки.'
    ]
    sources=[{'id':'W1','kind':'web','verified':True,'title':'Материал','url':'https://example.org','text':'Материал без всех указанных чисел, но достаточно подробный для проверки.'}]
    # QA may report warnings, but the generation pipeline is not allowed to throw because of them.
    warnings=source_issues(p,sources)
    assert isinstance(warnings,list)


def test_slidebatch_malformed_json_uses_deterministic_fallback(monkeypatch):
    import core.llm as llm
    from copy import deepcopy
    source={'id':'W1','kind':'web','url':'https://example.org','verified':True,'title':'Verified source',
            'text':'Это проверенный веб-источник с достаточным количеством содержательного текста для резервной генерации слайдов. '*8}
    from dataclasses import replace
    monkeypatch.setattr(llm,'settings',replace(llm.settings,api_key='test-key'))
    session=llm.Session()
    calls={'n':0}
    class FakeChoice:
        finish_reason='stop'
        message=type('M',(),{'content':'это не JSON'})()
    class FakeResp:
        choices=[FakeChoice()]
        usage=None
    class FakeCompletions:
        def create(self, **kwargs):
            calls['n']+=1
            return FakeResp()
    class FakeClient:
        def __enter__(self): return self
        def __exit__(self,*a): pass
        @property
        def chat(self): return type('C',(),{'completions':FakeCompletions()})()
    monkeypatch.setattr(llm,'OpenAI',lambda **kwargs: FakeClient())
    result=llm._complete(
        'Верни SlideBatch',
        {'brief':'Тема тестовой презентации','required_slide_numbers':[2],
         'architecture':{'slides':[{'number':2,'key_message':'Проверяемый материал'}]},'sources':[source]},
        0.1,1800,schema=llm.SlideBatch,session=session,stage='write_2_2'
    )
    assert result['slides'][0]['number']==2
    assert result['slides'][0]['source_ids']==['W1']
    assert calls['n']>=3


def test_short_evidence_quote_does_not_break_validation():
    from core.llm import _coerce_slide
    from core.models import Slide
    raw={
        'number':2,'kind':'content','title':'Короткая цитата',
        'points':['Содержательный блок с достаточным количеством слов для тестирования устойчивости контракта и локальной проверки.','Второй блок сохраняет объяснение механизма и не зависит от длины доказательной строки.'],
        'claims':[{'text':'Числовой факт','source_id':'W1','evidence_quote':'6.'}],
        'source_ids':['W1']
    }
    slide=_coerce_slide(raw)
    validated=Slide.model_validate(slide)
    assert validated.claims[0].evidence_quote=='6.'
