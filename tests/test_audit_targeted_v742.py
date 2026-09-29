from pathlib import Path
from tempfile import TemporaryDirectory
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from core.audit import audit_deck


def make_deck(path: Path):
    prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
    slide=prs.slides.add_slide(prs.slide_layouts[6])
    a=slide.shapes.add_textbox(Inches(1),Inches(1),Inches(5),Inches(1));a.name='generated:text'
    r=a.text_frame.paragraphs[0].add_run();r.text='TODO';r.font.name='Arial';r.font.size=Pt(28);r.font.color.rgb=RGBColor(255,0,0)
    b=slide.shapes.add_textbox(Inches(4.5),Inches(1.2),Inches(5),Inches(1));b.name='generated:text'
    b.text='Перекрывающийся текст для проверки аудита.'
    for rr in b.text_frame.paragraphs[0].runs:
        rr.font.name='Arial';rr.font.size=Pt(28)
    prs.save(path)


def test_expanded_audit_detects_placeholder_and_overlap():
    with TemporaryDirectory() as td:
        path=Path(td)/'audit.pptx';make_deck(path)
        out=audit_deck(path, {'slides':[{'number':1,'title':'TODO','subtitle':'','points':[],'left_points':[],'right_points':[],'left_title':'','right_title':'','quote':'','metrics':[],'timeline':[],'source_ids':[],'claims':[]}]})
        kinds={x['kind'] for x in out['issues']}
        assert 'placeholder_text' in kinds
        assert 'overlap' in kinds
